"""Reconcile the Jira TB backlog with the delivery workbook and validate the dependency graph.

Deterministic part of /plan-waves (see docs/workflow/README.md, section 3).

Usage:
    python scripts/workflow/reconcile_workbook.py <jira-export.json> [workbook.xlsx] [--json out.json]

<jira-export.json> is the raw result of a JQL search for `project = TB ORDER BY key ASC` with
fields summary, issuetype, status, priority, labels, parent, issuelinks, description,
customfield_10016, customfield_10020. Both shapes are accepted: the Atlassian MCP result
({"issues": {"nodes": [...]}}) and the REST API result ({"issues": [...]}).

Exits 1 when it finds a hard problem: unmapped or duplicate Blueprint IDs, missing stories,
link mismatches, cycles, dangling links, or a dependency on a later wave. Soft discrepancies
(sprint field, points field) are printed as warnings and do not fail the run.

Requires: openpyxl (pip install openpyxl).
"""

from __future__ import annotations

import argparse
import collections
import json
import re
import sys
from pathlib import Path

import openpyxl

PRIORITY_MAP = {"P0": "Highest", "P1": "High", "P2": "Medium", "P3": "Low"}
BLUEPRINT_RE = re.compile(r"Blueprint ID:\s*((?:US|EP)-\d+)")
STORY_RE = re.compile(r"US-\d+")
DEFAULT_WORKBOOK = "docs/TaskBuddy_Delivery_Workbook.xlsx"
DEFAULT_ACCEPTED = "docs/waves/accepted-deviations.txt"


def load_workbook(path: str) -> tuple[dict, dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    stories = {}
    for row in wb["Product Backlog"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        sid, epic, _feat, title, _story, _ac, pri, pts, wave, _ws, sprint, _phase, blocked, status = row[:14]
        stories[sid] = {
            "epic": epic,
            "title": str(title).strip(),
            "priority": pri,
            "points": int(pts or 0),
            "wave": int((re.findall(r"\d+", str(wave or "")) or ["0"])[0]),  # 0 = unassigned
            "sprint": int(sprint) if sprint else None,
            "deps": STORY_RE.findall(str(blocked or "")),
            "status": status,
        }
    epics = {row[0]: {"name": row[1]} for row in wb["Epic Backlog"].iter_rows(min_row=2, values_only=True) if row[0]}
    return stories, epics


def load_jira(path: str) -> dict:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    nodes = data["issues"]["nodes"] if isinstance(data.get("issues"), dict) else data["issues"]
    issues = {}
    for node in nodes:
        f = node["fields"]
        desc = str(f.get("description") or "")
        match = BLUEPRINT_RE.search(desc)
        blocked_by = [
            link["inwardIssue"]["key"]
            for link in f.get("issuelinks") or []
            if link["type"]["name"] == "Blocks" and "inwardIssue" in link
        ]
        issues[node["key"]] = {
            "blueprint": match.group(1) if match else None,
            "type": f["issuetype"]["name"],
            "title": f["summary"].strip(),
            "priority": (f.get("priority") or {}).get("name"),
            "labels": f.get("labels") or [],
            "parent": (f.get("parent") or {}).get("key"),
            "blocked_by": blocked_by,
            "status": f["status"]["name"],
            "sprints": [s["name"] for s in f.get("customfield_10020") or []],
            "points_field": f.get("customfield_10016"),
            "has_ac": bool(re.search(r"\\?\[[ xX]\\?\]", desc)),
        }
    return issues


def batches_within_waves(stories: dict) -> dict:
    """Batch index inside a wave: 0 if no same-wave dependency, else 1 + max batch of same-wave deps."""
    batch: dict[str, int] = {}

    def visit(sid: str) -> int:
        if sid not in batch:
            same = [d for d in stories[sid]["deps"] if d in stories and stories[d]["wave"] == stories[sid]["wave"]]
            batch[sid] = 0 if not same else 1 + max(visit(d) for d in same)
        return batch[sid]

    for sid in stories:
        visit(sid)
    return batch


def find_cycles(stories: dict) -> list[list[str]]:
    color: dict[str, int] = {}
    cycles = []

    def dfs(node: str, path: list[str]) -> None:
        color[node] = 1
        for dep in stories[node]["deps"]:
            if dep not in stories:
                continue
            if color.get(dep) == 1:
                cycles.append(path + [node, dep])
            elif not color.get(dep):
                dfs(dep, path + [node])
        color[node] = 2

    for sid in stories:
        if not color.get(sid):
            dfs(sid, [])
    return cycles


def critical_path(stories: dict) -> tuple[int, list[str]]:
    memo: dict[str, tuple[int, list[str]]] = {}

    def longest(sid: str) -> tuple[int, list[str]]:
        if sid not in memo:
            best: tuple[int, list[str]] = (0, [])
            for dep in stories[sid]["deps"]:
                if dep in stories and longest(dep)[0] > best[0]:
                    best = longest(dep)
            memo[sid] = (best[0] + stories[sid]["points"], best[1] + [sid])
        return memo[sid]

    return max((longest(s) for s in stories), key=lambda r: r[0])


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("jira_export")
    parser.add_argument("workbook", nargs="?", default=DEFAULT_WORKBOOK)
    parser.add_argument("--json", dest="json_out")
    parser.add_argument("--accept", default=DEFAULT_ACCEPTED,
                        help="file of accepted deviations: one error-message substring per line, # for comments")
    args = parser.parse_args()

    stories, epics = load_workbook(args.workbook)
    issues = load_jira(args.jira_export)
    errors: list[str] = []
    warnings: list[str] = []

    counts = collections.Counter(i["blueprint"] for i in issues.values() if i["blueprint"])
    errors += [f"duplicate Blueprint ID in Jira: {bp}" for bp, n in counts.items() if n > 1]
    errors += [f"Jira issue without Blueprint ID: {k}" for k, i in issues.items() if not i["blueprint"]]
    key_by_bp = {i["blueprint"]: k for k, i in issues.items() if i["blueprint"]}
    bp_by_key = {k: bp for bp, k in key_by_bp.items()}
    errors += [f"in workbook, missing in Jira: {s}" for s in [*stories, *epics] if s not in key_by_bp]
    errors += [f"in Jira, missing in workbook: {bp}" for bp in key_by_bp if bp not in stories and bp not in epics]

    for sid, s in stories.items():
        key = key_by_bp.get(sid)
        if not key:
            continue
        j = issues[key]
        tag = f"{sid}/{key}"
        if j["title"] != s["title"]:
            errors.append(f"{tag} title: workbook {s['title']!r} vs Jira {j['title']!r}")
        if PRIORITY_MAP.get(s["priority"]) != j["priority"]:
            errors.append(f"{tag} priority: workbook {s['priority']} vs Jira {j['priority']}")
        if f"wave-{s['wave']}" not in j["labels"]:
            errors.append(f"{tag} wave: workbook {s['wave']} vs labels {[l for l in j['labels'] if l.startswith('wave')]}")
        if j["parent"] != key_by_bp.get(s["epic"]):
            errors.append(f"{tag} parent: expected {key_by_bp.get(s['epic'])} got {j['parent']}")
        jira_deps = sorted(bp_by_key.get(b, b) for b in j["blocked_by"])
        if jira_deps != sorted(s["deps"]):
            errors.append(f"{tag} blocked-by: workbook {sorted(s['deps'])} vs Jira {jira_deps}")
        if not j["has_ac"]:
            errors.append(f"{tag} has no acceptance-criteria checklist in Jira")
        if s["sprint"] and j["sprints"] and j["sprints"] != [f"TB Sprint {s['sprint']}"]:
            warnings.append(f"{tag} sprint: workbook {s['sprint']} vs Jira {j['sprints']}")
        if j["points_field"] is None:
            warnings.append(f"{tag} points: field empty (workbook says {s['points']})")

    for sid, s in stories.items():
        for dep in s["deps"]:
            if dep not in stories:
                errors.append(f"dangling dependency {sid} -> {dep}")
            elif stories[dep]["wave"] > s["wave"]:
                errors.append(f"backward dependency: {sid} (wave {s['wave']}) needs {dep} (wave {stories[dep]['wave']})")
    cycles = find_cycles(stories)
    errors += [f"cycle: {' -> '.join(c)}" for c in cycles]
    if cycles:
        # Batches and critical path are undefined on a cyclic graph; report and stop.
        print(f"ERRORS ({len(errors)}) — dependency cycle found, plan cannot be computed")
        for e in errors:
            print("  " + e)
        return 1

    accepted_patterns = []
    if Path(args.accept).exists():
        accepted_patterns = [l.strip() for l in Path(args.accept).read_text(encoding="utf-8").splitlines()
                             if l.strip() and not l.lstrip().startswith("#")]
    accepted = [e for e in errors if any(p in e for p in accepted_patterns)]
    errors = [e for e in errors if e not in accepted]

    batch = batches_within_waves(stories)
    cp_points, cp = critical_path(stories)

    print(f"Jira issues: {len(issues)} | workbook stories: {len(stories)} epics: {len(epics)} | mapped: {len(key_by_bp)}")
    print(f"Critical path ({cp_points} pts): {' -> '.join(key_by_bp.get(s, s) for s in cp)}")
    for wave in sorted({s["wave"] for s in stories.values()}):
        ids = [s for s in stories if stories[s]["wave"] == wave]
        print(f"\nWave {wave}: {len(ids)} stories, {sum(stories[s]['points'] for s in ids)} pts")
        for b in sorted({batch[s] for s in ids}):
            print(f"  batch {chr(65 + b)}: " + ", ".join(key_by_bp.get(s, s) for s in ids if batch[s] == b))
    if warnings:
        kinds = collections.Counter(w.split(" ", 1)[1].split(":")[0] for w in warnings)
        print(f"\nWARNINGS ({len(warnings)}): " + ", ".join(f"{k} x{n}" for k, n in kinds.items()))
        for w in warnings[:10]:
            print("  " + w)
    if accepted:
        print(f"\nACCEPTED DEVIATIONS ({len(accepted)})")
        for a in accepted:
            print("  " + a)
    print(f"\nERRORS ({len(errors)})")
    for e in errors:
        print("  " + e)

    if args.json_out:
        Path(args.json_out).write_text(json.dumps(
            {"key_by_blueprint": key_by_bp, "stories": stories, "batches": batch,
             "critical_path": cp, "errors": errors, "warnings": warnings}, indent=1), encoding="utf-8")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
